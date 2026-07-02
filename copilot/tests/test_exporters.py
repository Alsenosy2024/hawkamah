"""Every exporter must produce a valid, non-empty file from Arabic markdown."""

import io
import zipfile

import pytest
from reportlab.pdfbase import pdfmetrics

from hawkama_copilot.exporters import export

MD = """# دليل الحوكمة

## الملخص التنفيذي
هذه فقرة عربية تجريبية مع مصطلح English ورقم 2026.

- بند أول
- بند ثانٍ

| البعد | النضج | الفجوة |
|---|---|---|
| الحوكمة | 3 | غياب ميثاق المجلس<br>وغياب لائحة العمل |
| المخاطر | 2 | لا يوجد سجل مخاطر & لا لجنة |
| الأداء | خط الأساس < المستهدف | تحسين |

> ملاحظة مرجعية.
"""


def _is_zip(data: bytes) -> bool:
    return zipfile.is_zipfile(io.BytesIO(data))


def test_markdown_and_txt():
    assert export(MD, "t", "md").data.decode("utf-8").startswith("# دليل")
    assert "دليل الحوكمة" in export(MD, "t", "txt").data.decode("utf-8")


def test_html_is_rtl():
    html = export(MD, "دليل", "html").data.decode("utf-8")
    assert 'dir="rtl"' in html
    assert "البعد" in html
    assert "<table" in html


def test_json_blocks():
    import json
    payload = json.loads(export(MD, "t", "json").data.decode("utf-8"))
    assert any(b["type"] == "table" for b in payload["blocks"])


MERMAID_MD = """# دليل بمخطط

## الهيكل
نص قبل المخطط.

```mermaid
graph TD
  A["مجلس الإدارة"] --> B["الإدارة التنفيذية"]
```
"""


def test_html_renders_mermaid_block_and_injects_runtime_once():
    html = export(MERMAID_MD, "دليل", "html").data.decode("utf-8")
    # the diagram is wrapped for the Mermaid runtime, not dumped as raw <pre><code>
    assert '<pre class="mermaid">' in html
    assert "graph TD" in html
    assert "<pre><code>graph TD" not in html
    # the runtime is injected exactly once, in the head
    assert html.count("mermaid.esm.min.mjs") == 1
    assert "mermaid.initialize" in html
    assert html.index("mermaid.esm.min.mjs") < html.index("<body")


def test_html_without_mermaid_omits_runtime():
    html = export(MD, "دليل", "html").data.decode("utf-8")
    assert "mermaid.esm.min.mjs" not in html


def test_docx_is_valid_ooxml():
    data = export(MD, "دليل الحوكمة", "docx", company="شركة").data
    assert _is_zip(data)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert "word/document.xml" in z.namelist()
        doc_xml = z.read("word/document.xml").decode("utf-8")
        assert "w:bidi" in doc_xml  # RTL applied


def test_docx_embeds_almarai_font():
    """P1/D5 — the docx exporter must EMBED the Almarai TTFs (mirrors the PDF
    exporter's vendored-font approach), not just name the font, so the document
    renders in-brand on a machine without Almarai installed."""
    from hawkama_copilot.exporters import docx_export

    assert (docx_export._FONT_DIR / "Almarai-Regular.ttf").is_file()
    assert (docx_export._FONT_DIR / "Almarai-Bold.ttf").is_file()

    data = export(MD, "دليل الحوكمة", "docx", company="شركة").data
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        # no duplicate/competing parts (the template already ships a fontTable)
        assert len(names) == len(set(names))
        assert names.count("word/fontTable.xml") == 1

        font_table = z.read("word/fontTable.xml").decode("utf-8")
        assert 'w:name="Almarai"' in font_table
        assert "w:embedRegular" in font_table and "w:embedBold" in font_table

        font_data_parts = [n for n in names if n.startswith("word/fonts/") and n.endswith(".fntdata")]
        assert len(font_data_parts) == 2  # regular + bold, non-empty TTF bytes
        for n in font_data_parts:
            assert len(z.read(n)) > 1000

        settings_xml = z.read("word/settings.xml").decode("utf-8")
        assert 'w:embedTrueTypeFonts w:val="true"' in settings_xml

        content_types = z.read("[Content_Types].xml").decode("utf-8")
        assert "application/x-font-ttf" in content_types

    # The saved file must still be a valid, re-openable docx.
    from docx import Document as _Document

    reopened = _Document(io.BytesIO(data))
    assert len(reopened.paragraphs) > 0


def test_xlsx_has_table_sheet():
    data = export(MD, "دليل", "xlsx").data
    assert _is_zip(data)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert len(wb.sheetnames) >= 2  # outline + at least one table sheet


def test_pptx_is_valid():
    data = export(MD, "دليل", "pptx", company="شركة").data
    assert _is_zip(data)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert any(n.startswith("ppt/slides/slide") for n in z.namelist())


def test_pdf_is_valid():
    data = export(MD, "دليل الحوكمة", "pdf", company="شركة").data
    assert data[:5] == b"%PDF-"
    assert len(data) > 1000


def test_pdf_uses_thmanyah_brand_font():
    """BE-2: the brand Thmanyah TTFs must be vendored inside the package (so the
    Cloud Run image has them) and be the PDF's first-choice regular/bold face."""
    from hawkama_copilot.exporters import pdf_export

    pkg = pdf_export._PKG_FONTS
    assert (pkg / "thmanyah-sans-regular.ttf").is_file()
    assert (pkg / "thmanyah-sans-bold.ttf").is_file()

    # Registration resolves the regular face to the vendored Thmanyah TTF.
    pdf_export._register_fonts()
    reg = pdfmetrics.getFont(pdf_export._FONT)
    assert "thmanyah" in str(reg.face.filename).lower()


def test_pdf_shaping_keeps_all_letters():
    """BE-2 RTL correctness: word-initial alef/dal must not be dropped.

    The brand/body fonts lack the isolated presentation forms U+FE8D/U+FEA9, so
    the reshaper must fall back to the base character whose glyph IS in the cmap —
    otherwise «دليل الحوكمة» loses its leading «د»/«ا»."""
    from hawkama_copilot.exporters.pdf_export import _ar, _register_fonts
    from reportlab.pdfbase import pdfmetrics

    _register_fonts()
    cmap = pdfmetrics.getFont("Body").face.charToGlyph
    for sample in ("دليل الحوكمة", "إدارة المخاطر", "الموارد البشرية"):
        shaped = _ar(sample)
        missing = [c for c in shaped if c != " " and ord(c) not in cmap]
        assert not missing, f"{sample!r} dropped glyphs: {[hex(ord(c)) for c in missing]}"


def test_unsupported_format_raises():
    with pytest.raises(ValueError):
        export(MD, "t", "rtf")
